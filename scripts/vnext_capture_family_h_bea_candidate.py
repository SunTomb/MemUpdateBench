"""Capture the bounded official BEA GDP revision trajectory without release claims."""
from __future__ import annotations

import argparse
from datetime import date
from io import BytesIO
import hashlib
import json
from pathlib import Path
import math
import re
import urllib.request

import openpyxl

SOURCE_GROUP_ID = "family-h-cross-domain-bea-gdp-2024q4"
SOURCE_DOCUMENT_ID = "bea-gdp-2024q4-revision-trajectory"
CAPTURE_SCHEMA = "memupdatebench.family-h.bea-source-capture.v1"
INDEX_SCHEMA = "memupdatebench.family-h.bea-source-capture-index.v1"
POLICY_URL = "https://www.bea.gov/about/policies-and-information"
SCHEDULE_URL = "https://www.bea.gov/news/schedule/full-2025"
USER_AGENT = "MemUpdateBench-Family-H-BEA-capture/1.0"

RELEASES = (
    (
        "advance", "2025-01-30", "https://www.bea.gov/sites/default/files/2025-01/gdp4q24-adv.xlsx",
        "https://www.bea.gov/news/2025/gross-domestic-product-fourth-quarter-and-year-2024-advance-estimate", 2.3,
    ),
    (
        "second", "2025-02-27", "https://www.bea.gov/sites/default/files/2025-02/gdp4q24-2nd.xlsx",
        "https://www.bea.gov/news/2025/gross-domestic-product-4th-quarter-and-year-2024-second-estimate", 2.3,
    ),
    (
        "third", "2025-03-27", "https://www.bea.gov/sites/default/files/2025-03/gdp4q24-3rd.xlsx",
        "https://www.bea.gov/sites/default/files/2025-03/gdp4q24-3rd.pdf", 2.4,
    ),
)
_RELEASE_BY_STAGE = {row[0]: row for row in RELEASES}


def canonical(value) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False) + "\n").encode("utf-8")


def sha(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def _text(value) -> str:
    return "" if value is None else " ".join(str(value).split())


def _header_text(value) -> str:
    return re.sub(r"\s+[a-z]$", "", _text(value), flags=re.IGNORECASE)


def _exact_release(stage: str, release_date: str, xlsx_url: str):
    _require(stage in _RELEASE_BY_STAGE, "unknown BEA estimate stage")
    expected = _RELEASE_BY_STAGE[stage]
    _require((stage, release_date, xlsx_url) == expected[:3], "BEA release identity mismatch")
    return expected


def parse_bea_xlsx(raw: bytes, stage: str, release_date: str, xlsx_url: str, page_url: str = "") -> dict:
    """Extract one exact Real GDP Q4 2024 annualized percentage from an XLSX."""
    expected = _exact_release(stage, release_date, xlsx_url)
    workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        matches = []
        for worksheet in workbook.worksheets:
            if worksheet.title != "Table 1":
                continue
            cells = {(cell.row, cell.column): _text(cell.value) for row in worksheet.iter_rows() for cell in row if cell.value is not None}
            real_rows = [position for position, text in cells.items() if text.casefold() == "real gdp"]
            bea_rows = [position for position, text in cells.items() if text.casefold() == "gross domestic product (gdp)"]
            _require(not (real_rows and bea_rows), "Real GDP row is ambiguous")
            row_candidates = real_rows or bea_rows
            _require(len(row_candidates) == 1, "Real GDP row is missing or ambiguous")
            quarter_columns = [position for position, text in cells.items() if _header_text(text).casefold() == "q4 2024"]
            if not quarter_columns:
                quarter_columns = [
                    (row, column) for (row, column), text in cells.items()
                    if _header_text(text).casefold() == "q4" and _header_text(cells.get((row - 1, column))).casefold() == "2024"
                ]
            _require(len(quarter_columns) == 1, "Q4 2024 column is missing or ambiguous")
            title_text = " ".join(cells.values()).casefold()
            unit_cells = [text for text in cells.values() if text.casefold() == "percent change from q3 to q4"]
            fixture_measure = len(unit_cells) == 1 and "percent change from q3 to q4" in title_text
            bea_measure = "percent change from preceding period" in title_text
            _require(fixture_measure or bea_measure, "required Q3-to-Q4 percent-change table missing")
            _require("annual rates" in title_text, "annualized rate basis missing")
            row, _ = row_candidates[0]
            _, column = quarter_columns[0]
            value = worksheet.cell(row=row, column=column).value
            _require(type(value) in {int, float} and not isinstance(value, bool) and math.isfinite(float(value)), "BEA value is not numeric")
            _require(round(float(value), 1) == float(value), "BEA value must retain only rounded one-decimal precision")
            matches.append((worksheet.title, row, column, float(value)))
        _require(len(matches) == 1, "BEA table worksheet is missing or ambiguous")
        worksheet, row, column, value = matches[0]
        _require(value == expected[4], "BEA value does not match the pinned release value")
        return {
            "record_id": f"bea-gdp-2024q4-{stage}",
            "source_group_id": SOURCE_GROUP_ID,
            "source_document_id": SOURCE_DOCUMENT_ID,
            "release_date": release_date,
            "estimate_stage": stage,
            "reference_period": "2024-Q4",
            "unit": "percent",
            "rate_basis": "annualized",
            "table_label": "Real GDP",
            "value": value,
            "xlsx_url": xlsx_url,
            "page_url": page_url,
            "table_url": xlsx_url,
            "source_file_sha256": sha(raw),
            "worksheet": worksheet,
            "row": row,
            "column": column,
            "source_anchor": {"worksheet": worksheet, "row": row, "column": column},
            "object_key": {
                "object_type": "macroeconomic_estimate", "namespace": "family_h", "entity": "us-gdp-2024q4",
                "attribute": "real_gdp_growth_annual_rate", "subkey": None,
            },
        }
    finally:
        workbook.close()


def fetch(url: str, output: Path) -> dict:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=30) as response:
        raw = response.read(32 * 1024 * 1024 + 1)
        if len(raw) > 32 * 1024 * 1024:
            raise ValueError("official BEA source exceeds bounded capture size")
        status = response.status
        content_type = response.headers.get("Content-Type")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(raw)
    return {"url": url, "http_status": status, "content_type": content_type, "bytes": len(raw), "sha256": sha(raw),
            "path": output.as_posix()}


def _write_capture(root: Path, manifest: dict) -> dict:
    root.mkdir(parents=True, exist_ok=True)
    (root / "capture_manifest.json").write_bytes(canonical(manifest))
    artifacts = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.name != "capture_index.json":
            artifacts.append({"path": path.relative_to(root).as_posix(), "bytes": path.stat().st_size, "sha256": sha(path.read_bytes())})
    index = {"schema": INDEX_SCHEMA, "status": manifest["status"], "artifacts": artifacts, "scientific_release_allowed": False}
    (root / "capture_index.json").write_bytes(canonical(index))
    return {"status": manifest["status"], "capture_manifest_sha256": sha((root / "capture_manifest.json").read_bytes()),
            "capture_index_sha256": sha((root / "capture_index.json").read_bytes()), "source_group_id": SOURCE_GROUP_ID}


def capture(output_root: Path) -> dict:
    output_root = Path(output_root).absolute()
    if output_root.exists():
        raise FileExistsError(output_root)
    records = []
    calls = []
    try:
        for stage, release_date, xlsx_url, page_url, _ in RELEASES:
            xlsx_call = fetch(xlsx_url, output_root / "raw" / f"{stage}.xlsx")
            xlsx_call["path"] = f"raw/{stage}.xlsx"
            page_suffix = ".pdf" if page_url.lower().endswith(".pdf") else ".html"
            page_path = output_root / "raw" / f"{stage}{page_suffix}"
            page_call = fetch(page_url, page_path)
            page_call["path"] = f"raw/{stage}{page_suffix}"
            calls.extend((xlsx_call, page_call))
            normalized = parse_bea_xlsx((output_root / "raw" / f"{stage}.xlsx").read_bytes(), stage, release_date, xlsx_url, page_url)
            records.append({"stage": stage, "release_date": release_date, "xlsx_url": xlsx_url, "page_url": page_url,
                            "xlsx": xlsx_call, "page": page_call, "normalized": normalized})
        policy = fetch(POLICY_URL, output_root / "raw" / "policy.html")
        policy["path"] = "raw/policy.html"
        schedule = fetch(SCHEDULE_URL, output_root / "raw" / "schedule.html")
        schedule["path"] = "raw/schedule.html"
    except Exception as exc:
        manifest = {
            "schema": CAPTURE_SCHEMA, "status": "BLOCKED_SOURCE_CAPTURE", "source_group_id": SOURCE_GROUP_ID,
            "source_document_id": SOURCE_DOCUMENT_ID, "records": [], "calls": calls,
            "policy_url": POLICY_URL, "schedule_url": SCHEDULE_URL, "policy_status": "POLICY_PENDING",
            "blocker": {"kind": "BLOCKED_SOURCE_CAPTURE", "detail": str(exc)},
            "formal_task_release": False, "scientific_release_allowed": False, "model_loads": 0, "generations": 0,
            "provider_calls": 0,
        }
        return _write_capture(output_root, manifest)
    manifest = {
        "schema": CAPTURE_SCHEMA, "status": "CAPTURED_PENDING_HUMAN_REVIEW", "source_group_id": SOURCE_GROUP_ID,
        "source_document_id": SOURCE_DOCUMENT_ID, "source_kind": "sequential_release_vintage", "source_type": "other",
        "domain": "macroeconomics", "language": "en", "revision_only": True, "independent_samples": False,
        "records": records, "calls": calls, "policy": {**policy, "url": POLICY_URL}, "schedule": {**schedule, "url": SCHEDULE_URL},
        "policy_status": "POLICY_PENDING", "formal_task_release": False, "scientific_release_allowed": False,
        "model_loads": 0, "generations": 0, "provider_calls": 0,
    }
    return _write_capture(output_root, manifest)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args(argv)
    print(json.dumps(capture(args.output_root), sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
